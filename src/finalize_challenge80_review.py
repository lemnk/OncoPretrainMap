"""Validate and release the completed Challenge-80 human review."""

from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data/validation/challenge80_caleb_blinded_review_corrected_v1.xlsx"
PRIVATE_KEY = ROOT / "data/validation/private/challenge80_key_v1.json"
FREEZE = ROOT / "data/validation/challenge80_freeze_v1.json"
RELEASED_KEY = ROOT / "data/validation/challenge80_released_key_v1.json"
ROW_COMPARISON = ROOT / "data/validation/challenge80_row_comparison_v1.csv"
SUMMARY = ROOT / "reports/challenge80_review_summary.json"
CONFUSION = ROOT / "reports/challenge80_confusion_matrix.csv"
LABELS = ["D0", "D1", "D2", "D3", "D4"]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> None:
    freeze = json.loads(FREEZE.read_text(encoding="utf-8"))
    key_path = PRIVATE_KEY if PRIVATE_KEY.exists() else RELEASED_KEY
    if sha256(key_path) != freeze["private_key_sha256"]:
        raise ValueError("Challenge-80 key does not match its prereview commitment")
    key = json.loads(key_path.read_text(encoding="utf-8"))
    key_by_id = {row["review_id"]: row for row in key}
    if len(key_by_id) != 80:
        raise ValueError("Expected 80 unique records in the Challenge-80 key")

    workbook = load_workbook(WORKBOOK, read_only=True, data_only=False)
    sheet = workbook["Blinded review"]
    headers = [sheet.cell(5, column).value for column in range(1, 17)]
    rows = [
        dict(zip(headers, [sheet.cell(row, column).value for column in range(1, 17)]))
        for row in range(6, 86)
    ]
    if len({row["Review ID"] for row in rows}) != 80:
        raise ValueError("Expected 80 unique completed review rows")

    comparison: list[dict[str, str]] = []
    confusion: Counter[tuple[str, str]] = Counter()
    for row in rows:
        review_id = row["Review ID"]
        if review_id not in key_by_id:
            raise ValueError(f"Unknown review ID: {review_id}")
        truth = key_by_id[review_id]
        initial_class = truth["initial_class"].split("_", 1)[0]
        reviewer_class = row["Reviewer class D0–D4"]
        if initial_class not in LABELS or reviewer_class not in LABELS:
            raise ValueError(f"Invalid class for {review_id}")
        if row["Reviewer name"] != "Caleb Yitna Ref":
            raise ValueError(f"Unexpected reviewer identity for {review_id}")
        if row["Blinded to registry answers?"] != "Yes":
            raise ValueError(f"Review was not recorded as blinded for {review_id}")
        if row["AI used for decisions?"] != "No":
            raise ValueError(f"AI-use field was not No for {review_id}")
        confusion[(initial_class, reviewer_class)] += 1
        comparison.append(
            {
                "review_id": review_id,
                "source_id": truth["source_id"],
                "benchmark_context": row["Benchmark context"],
                "model_checkpoint": row["Model/checkpoint"],
                "evaluation_dataset_or_task": row["Evaluation dataset/task"],
                "initial_class": initial_class,
                "reviewer_class": reviewer_class,
                "agreement": "Yes" if initial_class == reviewer_class else "No",
                "reviewer_evidence_grade": row["Evidence grade A–D"],
                "checkpoint_version_resolved": row["Checkpoint/version resolved?"],
                "reviewer_source_urls": row["Primary source URLs"],
                "reviewer_reason": row["Source passage / evidence and containment reasoning"],
                "reviewer_search_log": row["Independent search log (including negative search)"],
                "reviewer_name": row["Reviewer name"],
                "review_date": row["Review date (YYYY-MM-DD)"].date().isoformat(),
                "initially_blinded": row["Blinded to registry answers?"],
                "reviewer_used_ai": row["AI used for decisions?"],
            }
        )

    initial_counts = Counter(row["initial_class"] for row in comparison)
    reviewer_counts = Counter(row["reviewer_class"] for row in comparison)
    agreement = sum(row["agreement"] == "Yes" for row in comparison)
    observed = agreement / len(comparison)
    expected = sum(initial_counts[label] * reviewer_counts[label] for label in LABELS) / (80 * 80)
    kappa = (observed - expected) / (1 - expected) if expected < 1 else None
    d1_upgrades = sum(
        row["initial_class"] == "D1" and row["reviewer_class"] in {"D2", "D3", "D4"}
        for row in comparison
    )
    d2_downgrades = sum(
        row["initial_class"] == "D2" and row["reviewer_class"] == "D1"
        for row in comparison
    )

    RELEASED_KEY.write_text(json.dumps(key, indent=2) + "\n", encoding="utf-8")
    with ROW_COMPARISON.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(comparison[0]))
        writer.writeheader()
        writer.writerows(comparison)
    with CONFUSION.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["initial_class/reviewer_class", *LABELS, "per_class_agreement"])
        for initial in LABELS:
            values = [confusion[(initial, reviewer)] for reviewer in LABELS]
            denominator = initial_counts[initial]
            per_class = confusion[(initial, initial)] / denominator if denominator else "NA"
            writer.writerow([initial, *values, per_class])

    summary = {
        "reviewer": "Caleb Yitna Ref",
        "review_date": "2026-09-29",
        "sampled_relationships": 80,
        "sample_design": (
            "SHA-256-seeded, class-stratified challenge drawn from relationships absent from "
            "the two earlier review packets; 10 D0, 37 D1, 28 D2, and 5 D3."
        ),
        "exact_agreement_count": agreement,
        "exact_agreement_proportion": observed,
        "unweighted_cohens_kappa": kappa,
        "initial_class_counts": dict(sorted(initial_counts.items())),
        "reviewer_class_counts": dict(sorted(reviewer_counts.items())),
        "per_class_agreement": {
            label: confusion[(label, label)] / initial_counts[label] if initial_counts[label] else None
            for label in LABELS
        },
        "d1_upgraded_to_d2_d3_d4": d1_upgrades,
        "d2_downgraded_to_d1": d2_downgrades,
        "checkpoint_version_resolved": dict(
            sorted(Counter(row["checkpoint_version_resolved"] for row in comparison).items())
        ),
        "initially_blinded": True,
        "reviewer_used_ai": False,
        "interpretation": (
            "This is a rule-reliability challenge using existing benchmark sources. It is not a "
            "third independent benchmark, an estimate of registry-wide accuracy, or D4 validation."
        ),
        "completed_workbook": str(WORKBOOK.relative_to(ROOT)).replace("\\", "/"),
        "completed_workbook_sha256": sha256(WORKBOOK),
        "prereview_key_commitment_sha256": freeze["private_key_sha256"],
    }
    SUMMARY.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
